from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Read markdown file
with open('53_FEATURES_DEEP_AUDIT_REPORT.md', 'r', encoding='utf-8') as f:
    md_content = f.read()

# Create Word document
doc = Document()
doc.styles['Normal'].font.name = '宋体'
doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_heading(text, level):
    h = doc.add_heading('', level=level)
    run = h.add_run(text)
    run.font.size = Pt([18, 16, 14, 12, 11, 10][level-1])
    return h

def add_paragraph(text, style='Normal'):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.size = Pt(12)
    return p

def add_table(data, headers=None):
    table = doc.add_table(rows=1 if headers else 0, cols=len(data[0]) if data else 0)
    table.style = 'Table Grid'
    
    if headers:
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True
    
    for row_data in data:
        row = table.add_row()
        cells = row.cells
        for i, cell_text in enumerate(row_data):
            cells[i].text = cell_text

# Parse markdown and convert to docx
lines = md_content.split('\n')
i = 0
while i < len(lines):
    line = lines[i].strip()
    
    if line.startswith('# '):
        add_heading(line[2:], level=1)
    elif line.startswith('## '):
        add_heading(line[3:], level=2)
    elif line.startswith('### '):
        add_heading(line[4:], level=3)
    elif line.startswith('|'):
        # Table detection (simple parsing)
        table_data = []
        while i < len(lines) and '|' in lines[i]:
            parts = [p.strip() for p in lines[i].split('|')[1:-1]]
            if all(parts):
                table_data.append(parts)
            i += 1
        if table_data:
            add_table(table_data)
        continue
    elif line.startswith('- **') or line.startswith('* **'):
        p = doc.add_paragraph(line, style='List Bullet')
        continue
    elif line == '' or line is None:
        pass  # Skip empty lines
    else:
        # Code block detection
        if line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.space_before = Pt(6)
            font = p.add_run('\n'.join(code_lines)).font
            font.name = 'Consolas'
            font.size = Pt(9)
            continue
        elif line.startswith('> ') or line.startswith('「') or line.startswith('**') or line.startswith('---'):
            # Quote or separator
            p = doc.add_paragraph(line.lstrip('> ').replace('**', ''), style='Intense Quote')
        else:
            add_paragraph(line)
    
    i += 1

# Save
output_file = '53_功能模块深度审计报告.docx'
doc.save(output_file)
print(f'Done: {output_file} created')
